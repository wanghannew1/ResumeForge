# -*- coding: utf-8 -*-
"""
简历转换工具 —— 直接生成版（无需 word 模板 / 配置 excel）

与模板版（简历转换工具-解决附件图片后缀重复问题.py）的差别：
    模板版：原始数据表 + word模板 + 配置excel 三件套 -> docxtpl 渲染
    直接版：仅原始数据表 -> python-docx 程序化排版生成同样式 docx
对应关系全部内置在下方常量区；排版复刻原模板实测结构：
    A4、四边距15mm、宋体16pt居中标题 + 12行x8列合并表格 + 表后证件段落。
图片插入前的人脸/文字方向摆正复用主模块的 make_upright()。
"""

import logging
import os
import threading
import time
from importlib import import_module
from os import path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from PIL import Image

# ---------------------------------------------------------------- 常量区 ----
TITLE = '长春吉大附中实验学校招聘教师个人信息表'

TEXT_FIELDS = [                      # (原始列名, 标签)
    ('姓名（必填）', '姓名'),
    ('性别（必填）', '性别'),
    ('政治面貌（必填）', '政治面貌'),
    ('出生年月（必填）', '出生年月'),
    ('籍贯（必填）', '籍贯'),
    ('民族（必填）', '民族'),
    ('身份证号（必填）', '身份证号'),
    ('联系电话（必填）', '联系电话'),
    ('本科毕业学校（必填）', '本科毕业学校'),
    ('本科专业（必填）', '本科专业'),
    ('本科毕业时间（必填）', '毕业时间'),
    ('硕士毕业学校（必填）', '硕士毕业学校'),
    ('硕士专业（必填）', '硕士专业'),
    ('硕士毕业时间（必填）', '毕业时间'),
    ('博士毕业学校（必填）', '博士毕业学校'),
    ('博士专业（必填）', '博士专业'),
    ('博士毕业时间（必填）', '毕业时间'),
    ('参加工作年限（必填）', '参加工作年限'),
    ('应聘学科（必填）', '应聘学科'),
    ('教师资格证（必填）', '教师资格证'),
    ('邮箱（必填）', '邮箱'),
]
LONG_TEXT_FIELDS = [
    ('教育经历（起止时间、毕业院校、专业、学历、学位、培养方式等）（必填）',
     '学习经历\n(从高中填起)'),
    ('工作经历（起止时间、工作单位、岗位等）（必填）', '工作经历'),
    ('获奖情况（必填）', '获奖情况'),
]
DATE_COL = '提交时间（自动）'
COL_UNIT = '应聘单位（必填）'
PIC_SPECS = [                        # (原始列名, 最大张数, 宽mm, 高mm, 保PNG)
    ('小一寸照片（必填）', 1, 35, 50, False),
    ('身份证（正、反面）（必填）', 2, 80, 50, False),
    ('毕业证、学位证、岗位所需相关资格证（必填）', 6, 120, 180, False),
    ('个人承诺（必填）', 1, 10, 5, True),   # 实测输出即10x5mm小图章式
]
PROMISE_PARAS = [
    '本人认可并郑重承诺：',
    '1.本人所填写的个人信息及提交的应聘材料均真实有效，'
    '如有虚假，愿意承担由此引起的一切责任。',
    '2.本人在应聘此岗位过程中从未被吉林省彩虹人才开发咨询服务有限公司'
    '收取过报名费、安置费、培训费等任何费用。',
]

# 复用主模块的文件查找与方向摆正（模块名含中文/连字符，按名导入）
_main = import_module('简历转换工具-解决附件图片后缀重复问题')
find_files = _main.find_files
make_upright = _main.make_upright


def _g(x):
    """NaN/None -> ''。"""
    return '' if x is None or x != x else str(x)


def _col(df, want):
    """取列：先精确匹配列名，再包含匹配（应对表头微调）。返回 Series。"""
    if want in df.columns:
        return df[want]
    core = want.replace('（必填）', '')
    for c in df.columns:
        if core in str(c):
            return df[c]
    raise KeyError(f'原始表缺少列：{want}')


def _set_font(run, name='仿宋', size=10.5, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), name)


def _cell_text(cell, text, align='center'):
    """清空单元格写文本，\n 软换行。"""
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if align == 'center'
                   else WD_ALIGN_PARAGRAPH.LEFT)
    lines = str(text).split('\n')
    for k, line in enumerate(lines):
        r = p.add_run(line)
        _set_font(r)
        if k < len(lines) - 1:
            r.add_break()
    return p


def _insert_pic(paragraph, pic_abspath, w_mm, h_mm, keep_png,
                auto_orient, image_path, log):
    """图片经方向摆正后按原始长宽比等比缩放入槽位（不拉伸变形，
    贴槽位长边、不超出），透明图存PNG其余JPEG。"""
    src = Image.open(pic_abspath)
    if auto_orient:
        src, delta, _rs = make_upright(src, path.basename(pic_abspath),
                                       image_path, log)
        log(f'插入前摆正:{path.basename(pic_abspath)} 逆时针{delta}°')
    img_ratio = src.width / src.height
    slot_ratio = w_mm / h_mm
    if img_ratio >= slot_ratio:                  # 偏宽：宽度顶满槽位
        out_w, out_h = w_mm, round(w_mm / img_ratio, 2)
    else:                                        # 偏高：高度顶满槽位
        out_h, out_w = h_mm, round(h_mm * img_ratio, 2)
    temp_folder = path.join(image_path, 'temp_pic')
    if not path.exists(temp_folder):
        os.makedirs(temp_folder)
    suffix = 'png' if keep_png else 'jpeg'
    temp_abspath = path.join(temp_folder, f'{time.time_ns()}.{suffix}')
    if keep_png:
        src.save(temp_abspath, format='PNG')          # 保留透明通道
    else:
        src.convert('RGB').save(temp_abspath, format='JPEG')
    run = paragraph.add_run()
    run.add_picture(temp_abspath, width=Mm(out_w), height=Mm(out_h))


def build_resume_doc(row_df, pic_files, commit_date,
                     auto_orient, image_path, log):
    """程序化复刻模板版式生成单份简历。

    row_df: 该行的单行 DataFrame（列名=原始表列名）
    pic_files: {图片列名: [绝对路径,...]}
    """
    def v(key):
        return _g(_col(row_df, key).iloc[0])

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin = sec.bottom_margin = Mm(15)
    sec.left_margin = sec.right_margin = Mm(15)

    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_font(tp.add_run(TITLE), name='宋体', size=16, bold=True)

    t = doc.add_table(rows=12, cols=8)
    t.style = 'Table Grid'
    widths = [26, 30, 20, 24, 10, 10, 24, 36]           # 合计180mm
    for row in t.rows:
        for ci, c in enumerate(row.cells):
            c.width = Mm(widths[ci])

    # 照片格：第0~5行最后一列纵向合并居中
    photo_cell = t.cell(0, 7)
    for r in range(1, 6):
        photo_cell = photo_cell.merge(t.cell(r, 7))
    photo_cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    pc_p = photo_cell.paragraphs[0]
    pc_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for f in pic_files['小一寸照片（必填）']:
        _insert_pic(pc_p, f, 35, 50, False, auto_orient, image_path, log)

    # 行0~5：标签|值|标签|值|标签(跨2)|值
    trio_rows = [
        (0, '姓名（必填）', '性别（必填）', '政治面貌（必填）', '政治面貌'),
        (1, '出生年月（必填）', '籍贯（必填）', '民族（必填）', '民族'),
        (3, '本科毕业学校（必填）', '本科专业（必填）',
         '本科毕业时间（必填）', '毕业时间'),
        (4, '硕士毕业学校（必填）', '硕士专业（必填）',
         '硕士毕业时间（必填）', '毕业时间'),
        (5, '博士毕业学校（必填）', '博士专业（必填）',
         '博士毕业时间（必填）', '毕业时间'),
    ]
    for r, k1, k2, k3, l3 in trio_rows:
        _cell_text(t.cell(r, 0), dict(TEXT_FIELDS)[k1])
        _cell_text(t.cell(r, 1), v(k1))
        _cell_text(t.cell(r, 2), dict(TEXT_FIELDS)[k2])
        _cell_text(t.cell(r, 3), v(k2))
        _cell_text(t.cell(r, 4).merge(t.cell(r, 5)), l3)
        _cell_text(t.cell(r, 6), v(k3))

    # 行2：身份证号 值跨4 | 联系电话 | 值
    _cell_text(t.cell(2, 0), '身份证号')
    _cell_text(t.cell(2, 1).merge(t.cell(2, 4)), v('身份证号（必填）'))
    _cell_text(t.cell(2, 5), '联系电话')
    _cell_text(t.cell(2, 6), v('联系电话（必填）'))

    # 行6/7：跨度 1|3|2|2
    _cell_text(t.cell(6, 0), '参加工作年限')
    _cell_text(t.cell(6, 1).merge(t.cell(6, 3)), v('参加工作年限（必填）'))
    _cell_text(t.cell(6, 4).merge(t.cell(6, 5)), '应聘学科')
    _cell_text(t.cell(6, 6).merge(t.cell(6, 7)), v('应聘学科（必填）'))
    _cell_text(t.cell(7, 0), '教师资格证')
    _cell_text(t.cell(7, 1).merge(t.cell(7, 3)), v('教师资格证（必填）'))
    _cell_text(t.cell(7, 4).merge(t.cell(7, 5)), '邮箱')
    _cell_text(t.cell(7, 6).merge(t.cell(7, 7)), v('邮箱（必填）'))

    # 行8~10 长文本：标签 | 值跨7（左对齐）
    for r, (key, label) in zip((8, 9, 10), LONG_TEXT_FIELDS):
        _cell_text(t.cell(r, 0), label)
        _cell_text(t.cell(r, 1).merge(t.cell(r, 7)), v(key), align='left')

    # 行11 个人承诺：固定文案 + 签名图(10x5mm PNG) + 提交日期
    _cell_text(t.cell(11, 0), '个人承诺')
    pc = t.cell(11, 1).merge(t.cell(11, 7))
    pc.text = ''
    for k, line in enumerate(PROMISE_PARAS):
        p = pc.paragraphs[0] if k == 0 else pc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _set_font(p.add_run(line))
    sig_p = pc.add_paragraph()
    sig_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _set_font(sig_p.add_run('               本人签名： '))
    for f in pic_files['个人承诺（必填）']:
        _insert_pic(sig_p, f, 10, 5, True, auto_orient, image_path, log)
    _set_font(sig_p.add_run('                    ' + commit_date))

    # 表后：身份证正反面同行居中；证书逐张一段（上限6张，模板版只有4槽位）
    p_ids = doc.add_paragraph()
    p_ids.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for f in pic_files['身份证（正、反面）（必填）']:
        _insert_pic(p_ids, f, 80, 50, False, auto_orient, image_path, log)
    for f in pic_files['毕业证、学位证、岗位所需相关资格证（必填）']:
        p_d = doc.add_paragraph()
        p_d.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _insert_pic(p_d, f, 120, 180, False, auto_orient, image_path, log)
    return doc


def xlsx2docx_resume_direct(raw_xlsx, start_n, end_n, image_path, docx_path,
                            auto_orient=True):
    """直接模式入口：仅原始数据表 -> word简历。

    raw_xlsx 可传 xlsx 路径，也可直接传已读好的 DataFrame(dtype=str)——
    用于图片不在单元格引用而需外部预处理注入的场景（如从xlsx内嵌媒体提取）。
    列名兼容：_col() 会剥离（必填）后缀做包含匹配。
    """
    import pandas as pd
    start_time = time.time()
    print(f"{time.strftime('%X')}[直接模式]开始导出")

    def log(msg):
        print(msg)
        logging.info(msg)

    if isinstance(raw_xlsx, pd.DataFrame):
        data = raw_xlsx.copy()
    else:
        data = pd.read_excel(raw_xlsx, dtype=str)
    total = data.shape[0]
    if start_n == 0 and end_n == 0:
        n1, n2 = 0, total
    else:
        n1 = start_n - 2 if start_n >= 2 else 0
        n2 = end_n - 1 if end_n >= 2 else 0

    for i in range(n1, n2):
        row_df = data.iloc[[i]]
        pic_files = {}
        for col, max_n, _w, _h, _png in PIC_SPECS:
            raw_names = _g(_col(data, col).iloc[i])
            names = [n.replace('/', '-').replace(':', '-')
                     for n in raw_names.split('\n') if n]
            found = []
            for n in names[:max_n]:
                hit = find_files(image_path, n)
                if hit:
                    found.append(hit)
                else:
                    log(f'[警告]找不到图片:{n}')
            pic_files[col] = found

        commit_raw = _g(_col(data, DATE_COL).iloc[i])
        commit_date = commit_raw.split()[0] if commit_raw else ''

        doc = build_resume_doc(row_df, pic_files, commit_date,
                               auto_orient, image_path, log)

        j = _g(_col(data, COL_UNIT).iloc[i])
        c = _g(_col(data, '姓名（必填）').iloc[i])
        n = _g(_col(data, '性别（必填）').iloc[i])
        fname = f'{j}_第{i + 2}行_{c}_{n}_{commit_date}.docx'
        doc.save(path.join(docx_path, fname))
        print(f'已导出第{i + 1}份简历（第{i + 2}行），姓名：{c}')

    print(f"{time.strftime('%X')}[直接模式]导出结束，"
          f"用时{round(time.time() - start_time, 2)}秒")


def start_thread_direct(raw_xlsx, start_n, end_n, image_path, docx_path,
                        auto_orient=True):
    thread = threading.Thread(
        target=xlsx2docx_resume_direct,
        args=(raw_xlsx, start_n, end_n, image_path, docx_path, auto_orient))
    thread.setDaemon(True)
    thread.start()
